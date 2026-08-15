"""Coletor auditável RAIS 2024: descoberta, raws validados e publicação atômica."""
from __future__ import annotations
import csv, ftplib, hashlib, json, mimetypes, re, shutil, tempfile, zipfile
from dataclasses import dataclass
from datetime import UTC,datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import HTTPError,URLError
from urllib.parse import urljoin,urlparse
from urllib.request import Request,build_opener
from ice_sul.extract.contracts import CollectionResult,CollectionStatus
from ice_sul.extract.http import TracingRedirect
from ice_sul.extract.registry import register_collector
from ice_sul.transform.rais import UF_PREFIX,load_municipality_map,safe_member,transform_archive

YEAR=2024; OFFICIAL_HOST="ftp.mtps.gov.br"; OFFICIAL_DIR="/pdet/microdados/RAIS/2024"; OFFICIAL_ROOT=f"https://{OFFICIAL_HOST}{OFFICIAL_DIR}"
PORTAL_HOST="www.gov.br"; IBGE_HOST="servicodados.ibge.gov.br"
PORTAL_PAGE="https://www.gov.br/trabalho-e-emprego/pt-br/acesso-a-informacao/acoes-e-programas/programas-projetos-acoes-obras-e-atividades/estatisticas-trabalho/rais/rais-2024/rais-2024-1"
DE_PARA_URL=PORTAL_PAGE+"/de-para-microdados.xlsx/@@download/file"
IBGE_URL="https://servicodados.ibge.gov.br/api/v1/localidades/municipios?orderBy=id"
UF_CODES=tuple(sorted(UF_PREFIX.values()))
XLSX_MIME="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
REQUIRED_FIELDS={"cnae20classecódigo","indvínculoativo3112código","naturezajurídicacódigo"}
MUNICIPAL_FIELDS={"municípiotrabcódigo","municípiocódigo"}

@dataclass(frozen=True)
class RaisArchive:
 uf:str; url:str; year:int=YEAR
 @property
 def filename(self)->str:return self.url.rstrip("/").rsplit("/",1)[-1]

class OfficialRoutesUnavailable(OSError):
 def __init__(self,attempts): super().__init__("listagens oficiais HTTPS e FTP indisponíveis"); self.attempts=attempts
class _Links(HTMLParser):
 def __init__(self):super().__init__();self.links=[]
 def handle_starttag(self,tag,attrs):
  if tag.lower()=="a" and (href:=dict(attrs).get("href")):self.links.append(href)

def sha256_file(path:Path,chunk_size:int=8*1024*1024)->str:
 digest=hashlib.sha256()
 with path.open("rb") as stream:
  while chunk:=stream.read(chunk_size):digest.update(chunk)
 return digest.hexdigest()

def content_type(path:Path)->str:
 return {".7z":"application/x-7z-compressed",".xlsx":XLSX_MIME,".json":"application/json"}.get(path.suffix.lower(),mimetypes.guess_type(path.name)[0] or "application/octet-stream")

def _manifest(*,url,method,path,uf="BR",status=None)->dict:
 return {"fonte":"MTE/PDET RAIS","url":url,"metodo":method,"parametros":{},"periodo":str(YEAR),"data_hora_utc":datetime.now(UTC).isoformat(),"status_http":status,"url_final":url,"redirects":[],"content_type":content_type(path),"uf":uf,"tamanho_transferido":0,"tamanho_persistido":path.stat().st_size if path.exists() else None,"sha256":sha256_file(path) if path.exists() else None,"arquivo":str(path),"licenca":"não informada na distribuição consultada","versao_snapshot":f"RAIS {YEAR}","validacao":"pendente","status":"downloaded"}

def validate_xlsx(path:Path)->dict:
 if not zipfile.is_zipfile(path): raise ValueError(f"XLSX inválido/não OOXML: {path}")
 try:
  from openpyxl import load_workbook
  workbook=load_workbook(path,read_only=True,data_only=True)
  try:
   sheets=workbook.sheetnames
   if "VINC_PUB" not in sheets: raise ValueError("workbook sem aba VINC_PUB")
   sheet=workbook["VINC_PUB"]
   if str(sheet["A1"].value).strip().lower() != "de" or str(sheet["B1"].value).strip().lower() != "para":
    raise ValueError("VINC_PUB sem estrutura De/Para oficial")
   target_values={str(row[1].value).strip().lower() for row in sheet.iter_rows(min_col=1,max_col=2) if row[1].value is not None}
   missing=REQUIRED_FIELDS-target_values
   if missing or not MUNICIPAL_FIELDS&target_values: raise ValueError(f"VINC_PUB sem campos obrigatórios: {sorted(missing | ({'campo_municipal'} if not MUNICIPAL_FIELDS&target_values else set()))}")
   confirmed=sorted(REQUIRED_FIELDS | (MUNICIPAL_FIELDS&target_values))
  finally: workbook.close()
 except (OSError,zipfile.BadZipFile,KeyError) as exc: raise ValueError(f"workbook XLSX corrompido: {exc}") from exc
 return {"tipo":"XLSX OOXML funcional","abas":sheets,"campos_confirmados":confirmed,"tamanho":path.stat().st_size,"sha256":sha256_file(path)}

def _validate_open_archive(seven)->dict:
 names=seven.getnames()
 if not names: raise ValueError("7-Zip vazio")
 for name in names:safe_member(name)
 candidates=[n for n in names if n.lower().endswith(".comt")]
 if len(candidates)!=1: raise ValueError(f"esperado exatamente um .comt de Vínculos; encontrados: {candidates}")
 bad_member=seven.testzip()
 if bad_member is not None: raise ValueError(f"CRC inválido no membro: {bad_member}")
 crc_result=seven.test()
 if crc_result is False: raise ValueError("teste CRC do 7-Zip falhou")
 return {"membros":names,"membro_comt":candidates[0],"crc_result":crc_result,"testzip":bad_member}

def validate_7z(path:Path)->dict:
 import py7zr
 if not py7zr.is_7zfile(path): raise ValueError(f"assinatura 7-Zip inválida: {path}")
 try:
  with py7zr.SevenZipFile(path) as seven: info=_validate_open_archive(seven)
  # Extração/leitura é obrigatória inclusive quando CRC é ausente.
  with tempfile.TemporaryDirectory(prefix="sit-rais-validate-") as d:
   with py7zr.SevenZipFile(path) as seven: seven.extract(path=d,targets=[info["membro_comt"]])
   member=Path(d)/safe_member(info["membro_comt"])
   with member.open("rb") as stream:
    if not stream.read(1): raise ValueError("membro .comt vazio/não extraível")
 except ValueError: raise
 except Exception as exc: raise ValueError(f"7-Zip corrompido/truncado: {exc}") from exc
 return {"tipo":"7-Zip","membros":info["membros"],"membro_comt":info["membro_comt"],"crc_result":info["crc_result"],"tamanho":path.stat().st_size,"sha256":sha256_file(path)}

def reused_manifest(path:Path,*,url:str,uf:str,validation:dict)->dict:
 entry=_manifest(url=url,method="REUSE",path=path,uf=uf); entry.update(status="reused",reutilizado=True,validacao=validation)
 return entry

def _download_validated(url:str,destination:Path,*,uf:str,validator,timeout=180)->dict:
 if destination.exists():raise FileExistsError(f"raw imutável já existe: {destination}")
 parsed=urlparse(url)
 if parsed.hostname not in {OFFICIAL_HOST,PORTAL_HOST,IBGE_HOST}:raise ValueError(f"host externo proibido: {parsed.hostname}")
 destination.parent.mkdir(parents=True,exist_ok=True); part=destination.with_name(destination.name+".part"); part.unlink(missing_ok=True)
 redirect=TracingRedirect(); opener=build_opener(redirect); digest=hashlib.sha256(); size=0
 try:
  with opener.open(Request(url,headers={"User-Agent":"SIT-mvp-demo-2026/1.0"}),timeout=timeout) as response,part.open("xb") as output:
   final=response.url
   if urlparse(final).hostname not in {OFFICIAL_HOST,PORTAL_HOST,IBGE_HOST}:raise ValueError(f"redirect para host externo: {final}")
   while chunk:=response.read(8*1024*1024):output.write(chunk);digest.update(chunk);size+=len(chunk)
   status=response.status; ctype=response.headers.get("Content-Type")
  if not size:raise ValueError("download vazio")
  validation=validator(part); part.replace(destination)
 except Exception:part.unlink(missing_ok=True);raise
 entry=_manifest(url=url,method="GET",path=destination,uf=uf,status=status)
 entry.update(url_final=final,redirects=redirect.history,content_type=ctype,tamanho_transferido=size,sha256=digest.hexdigest(),validacao=validation)
 return entry

def view_to_download(url:str)->str:
 parsed=urlparse(url)
 if parsed.hostname!=PORTAL_HOST:raise ValueError("link De-Para fora do domínio oficial")
 if re.search(r"\.xlsx/view/?$",parsed.path,re.I): return url[:url.lower().rfind("/view")]+"/@@download/file"
 if "@@download/file" in parsed.path:return url
 raise ValueError(f"link De-Para inesperado: {url}")

def discover_de_para_url(page_html:str)->str:
 parser=_Links();parser.feed(page_html)
 candidates=[]
 for link in parser.links:
  absolute=urljoin(PORTAL_PAGE+"/",link)
  if "de-para" in absolute.lower() and (".xlsx/view" in absolute.lower() or "@@download/file" in absolute.lower()):candidates.append(view_to_download(absolute))
 if len(set(candidates))!=1:raise ValueError(f"esperado um link De-Para oficial; encontrados {sorted(set(candidates))}")
 return candidates[0]

def download_de_para(destination:Path)->tuple[dict,list[dict]]:
 attempts=[]
 try:return _download_validated(DE_PARA_URL,destination,uf="BR",validator=validate_xlsx),attempts
 except (HTTPError,URLError,TimeoutError) as exc:attempts.append({"rota":"direta","url":DE_PARA_URL,"erro":f"{type(exc).__name__}: {exc}"})
 redirect=TracingRedirect(); opener=build_opener(redirect)
 try:
  with opener.open(Request(PORTAL_PAGE,headers={"User-Agent":"SIT-mvp-demo-2026/1.0"}),timeout=60) as response:
   if urlparse(response.url).hostname!=PORTAL_HOST:raise ValueError("redirect externo na página De-Para")
   url=discover_de_para_url(response.read().decode("utf-8","replace"))
  entry=_download_validated(url,destination,uf="BR",validator=validate_xlsx); attempts.append({"rota":"pagina","url":PORTAL_PAGE,"status_http":response.status,"redirects":redirect.history})
  return entry,attempts
 except (HTTPError,URLError,TimeoutError) as exc:attempts.append({"rota":"pagina","url":PORTAL_PAGE,"erro":f"{type(exc).__name__}: {exc}"});raise OfficialRoutesUnavailable(attempts) from exc

def parse_archive_listing(items:list[str],*,source_url:str)->list[RaisArchive]:
 parsed_source=urlparse(source_url)
 if parsed_source.hostname!=OFFICIAL_HOST:raise ValueError(f"host de listagem externo: {source_url}")
 found={}
 for raw in items:
  parsed=urlparse(raw); name=Path(parsed.path or raw).name
  if parsed.hostname and parsed.hostname!=OFFICIAL_HOST:raise ValueError(f"URL externa na listagem: {raw}")
  if not name.lower().endswith(".7z"):continue
  if "ESTAB" in name.upper():continue
  match=re.fullmatch(r"RAIS_VINC_PUB_([A-Z]{2})(?:_2024)?\.7z",name,re.I)
  if not match:raise ValueError(f"nome .7z inesperado na listagem RAIS: {name}")
  uf=match.group(1).upper()
  if uf not in UF_CODES:raise ValueError(f"UF inesperada: {uf}")
  if uf in found:raise ValueError(f"UF duplicada na listagem: {uf}")
  found[uf]=RaisArchive(uf,urljoin(OFFICIAL_ROOT+"/",name))
 missing=sorted(set(UF_CODES)-found.keys())
 if missing:raise ValueError(f"listagem oficial incompleta; UFs ausentes: {missing}")
 return [found[uf] for uf in UF_CODES]
archives_from_listing=lambda links:parse_archive_listing(links,source_url=OFFICIAL_ROOT+"/")

def _https_listing()->tuple[list[str],dict]:
 redirect=TracingRedirect();opener=build_opener(redirect)
 with opener.open(Request(OFFICIAL_ROOT+"/",headers={"User-Agent":"SIT-mvp-demo-2026/1.0"}),timeout=180) as response:
  if urlparse(response.url).hostname!=OFFICIAL_HOST:raise ValueError("redirect externo na listagem HTTPS")
  body=response.read(); parser=_Links();parser.feed(body.decode("utf-8","replace"))
  return parser.links,{"metodo":"HTTPS","status":"success","status_http":response.status,"url_final":response.url,"redirects":redirect.history,"listagem_sha256":hashlib.sha256(body).hexdigest(),"itens":len(parser.links)}

def _ftp_listing()->tuple[list[str],list[dict]]:
 attempts=[]
 with ftplib.FTP() as ftp:
  ftp.connect(OFFICIAL_HOST,21,timeout=180);ftp.login();ftp.cwd(OFFICIAL_DIR)
  try:
   entries=list(ftp.mlsd()); names=[name for name,facts in entries if facts.get("type") in {"file",None}];attempts.append({"metodo":"FTP MLSD","status":"success","itens":len(names)})
  except ftplib.all_errors as exc:
   attempts.append({"metodo":"FTP MLSD","status":"failed","erro":f"{type(exc).__name__}: {exc}"});names=ftp.nlst();attempts.append({"metodo":"FTP NLST","status":"success","itens":len(names)})
 return names,attempts

def discover_official_archives()->tuple[list[RaisArchive],list[dict]]:
 attempts=[]
 try:items,evidence=_https_listing();attempts.append(evidence);return parse_archive_listing(items,source_url=OFFICIAL_ROOT+"/"),attempts
 except (HTTPError,URLError,TimeoutError,OSError) as exc:attempts.append({"metodo":"HTTPS","status":"failed","erro":f"{type(exc).__name__}: {exc}"})
 try:
  items,ftp_attempts=_ftp_listing();attempts.extend(ftp_attempts);raw="\n".join(items).encode();attempts.append({"metodo":"FTP listagem usada","status":"success","listagem_sha256":hashlib.sha256(raw).hexdigest()})
  return parse_archive_listing(items,source_url=f"ftp://{OFFICIAL_HOST}{OFFICIAL_DIR}/"),attempts
 except ftplib.all_errors as exc:attempts.append({"metodo":"FTP","status":"failed","erro":f"{type(exc).__name__}: {exc}"});raise OfficialRoutesUnavailable(attempts) from exc

def _download_ftp_validated(spec:RaisArchive,destination:Path)->dict:
 destination.parent.mkdir(parents=True,exist_ok=True);part=destination.with_name(destination.name+".part");part.unlink(missing_ok=True);digest=hashlib.sha256();size=0
 try:
  with ftplib.FTP() as ftp:
   ftp.connect(OFFICIAL_HOST,21,timeout=180);ftp.login()
   with part.open("xb") as output:
    def receive(chunk):
     nonlocal size
     output.write(chunk);digest.update(chunk);size+=len(chunk)
    reply=ftp.retrbinary(f"RETR {OFFICIAL_DIR}/{spec.filename}",receive,blocksize=8*1024*1024)
  if not size:raise ValueError("download FTP vazio")
  validation=validate_7z(part);part.replace(destination)
 except Exception:part.unlink(missing_ok=True);raise
 entry=_manifest(url=f"ftp://{OFFICIAL_HOST}{OFFICIAL_DIR}/{spec.filename}",method="FTP RETR",path=destination,uf=spec.uf)
 entry.update(resposta_ftp=reply,tamanho_transferido=size,sha256=digest.hexdigest(),validacao=validation);return entry

def _download_archive(spec:RaisArchive,path:Path)->dict:
 try:return _download_validated(spec.url,path,uf=spec.uf,validator=validate_7z)
 except (HTTPError,URLError,TimeoutError,OSError) as exc:
  entry=_download_ftp_validated(spec,path);entry["tentativas_anteriores"]=[{"metodo":"HTTPS","erro":f"{type(exc).__name__}: {exc}"}];return entry

def _write_csv(path:Path,rows:list[dict]):
 with path.open("w",newline="",encoding="utf-8") as f:writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)

@dataclass
class RaisCollector:
 destination:Path=Path("data/raw/rais/2024");interim:Path=Path("data/interim/rais/2024");quality:Path=Path("reports/quality/mvp-demo-2026/rais");south_reference:Path=Path("data/processed/2026/municipios.csv");archives:list[RaisArchive]|None=None;source:str="rais"
 def collect(self)->CollectionResult:
  entries=[];artifacts=[];warnings=[];final_private=self.interim/"emprego_privado_municipal.csv";final_diag=self.interim/"mer_diag_01.csv";qa_path=self.quality/"qa.json"
  for p in (final_private,final_diag,qa_path):p.unlink(missing_ok=True)
  self.destination.mkdir(parents=True,exist_ok=True);self.quality.mkdir(parents=True,exist_ok=True)
  def fail(status,error,extra=None):
   qa={"ufs_esperadas":list(UF_CODES),"ufs_processadas":[],"arquivos_esperados":[],"arquivos_processados":[],"coverage_complete":False,"municipios_estoque_positivo":0,"vinculos_privados_elegiveis":0,"codigos_nao_ligados":{},"reconciliacao_municipal_ok":False,"reconciliacao_territorial_ok":False,"valores_invalidos":1 if "inválid" in str(error) else 0,"erro":str(error),**(extra or {})};qa_path.write_text(json.dumps(qa,ensure_ascii=False,indent=2)+"\n",encoding="utf-8");return CollectionResult(self.source,status,artifacts,["MER-02","MER-DIAG-01"],entries,warnings,[str(error)])
  try:
   ibge_path=self.destination/"municipios-ibge.json"
   if not ibge_path.exists():
    # JSON não é binário de alto risco; ainda assim baixa para .part e valida antes de promover.
    def validate_json(p):
     rows=json.loads(p.read_text());
     if not isinstance(rows,list) or not rows:raise ValueError("cadastro IBGE nacional inválido")
     return {"tipo":"JSON","registros":len(rows),"sha256":sha256_file(p)}
    entries.append(_download_validated(IBGE_URL,ibge_path,uf="BR",validator=validate_json))
   else:
    rows=json.loads(ibge_path.read_text());
    if not isinstance(rows,list) or not rows:raise ValueError("cadastro IBGE nacional inválido")
    entries.append(reused_manifest(ibge_path,url=IBGE_URL,uf="BR",validation={"tipo":"JSON","registros":len(rows),"sha256":sha256_file(ibge_path)}))
   ibge_rows=json.loads(ibge_path.read_text()); municipality_map={}
   for item in ibge_rows:
    code=str(item["id"])
    if not re.fullmatch(r"\d{7}",code) or code[:6] in municipality_map:raise ValueError(f"cadastro IBGE inválido/ambíguo: {code}")
    municipality_map[code[:6]]=code;municipality_map[code]=code
   south_map=load_municipality_map(self.south_reference);south_ids=sorted({v for k,v in south_map.items() if len(k)==6})
   if len(south_ids)!=1191:raise ValueError(f"universo Sul exige 1.191 municípios; obtidos {len(south_ids)}")
   depara=self.destination/"De-Para Microdados.xlsx"
   if depara.exists():entries.append(reused_manifest(depara,url=DE_PARA_URL,uf="BR",validation=validate_xlsx(depara)))
   else:entry,attempts=download_de_para(depara);entries.extend(attempts);entries.append(entry)
   artifacts.append(str(depara))
   specs=self.archives
   if specs is None:specs,attempts=discover_official_archives();entries.extend(attempts)
   # Injeção de teste passa pelo mesmo portão nacional.
   if len(specs)!=27 or {s.uf for s in specs}!=set(UF_CODES) or len({s.uf for s in specs})!=27:raise ValueError("cobertura de arquivos exige exatamente as 27 UFs únicas")
   results=[];processed=[]
   for spec in specs:
    if spec.year!=YEAR or urlparse(spec.url).hostname not in {OFFICIAL_HOST,None}:raise ValueError(f"arquivo fora da edição/host oficial: {spec}")
    archive=self.destination/spec.filename
    if archive.exists():entries.append(reused_manifest(archive,url=spec.url,uf=spec.uf,validation=validate_7z(archive)))
    else:entries.append(_download_archive(spec,archive))
    artifacts.append(str(archive));result=transform_archive(archive,file_uf=spec.uf,municipality_map=municipality_map,year=YEAR,south_municipalities=[c for c in south_ids if UF_PREFIX.get(c[:2])==spec.uf]);results.append(result);processed.append(spec)
   private=[row for r in results for row in r.private_employment];diagnostic=[row for r in results for row in r.diversification]
   if len(private)!=len({r["municipio_id"] for r in private}):raise ValueError("município duplicado entre arquivos de UF")
   unmatched={s.uf:r.quality["codigos_nao_ligados"] for s,r in zip(specs,results) if r.quality["codigos_nao_ligados"]}
   reconciliations=all(r.quality["reconciliacao_ok"] and r.quality["reconciliacao_territorial_ok"] for r in results)
   if unmatched:raise ValueError(f"códigos municipais ativos não ligados: {unmatched}")
   if len(diagnostic)!=1191 or len({r["municipio_id"] for r in diagnostic})!=1191:raise ValueError("MER-DIAG-01 não contém 1.191 municípios únicos")
   if not private or not reconciliations:raise ValueError("produto nacional vazio ou reconciliação falhou")
   qa={"ano":YEAR,"ufs_esperadas":list(UF_CODES),"ufs_processadas":[s.uf for s in processed],"arquivos_esperados":[s.filename for s in specs],"arquivos_processados":[s.filename for s in processed],"coverage_complete":True,"municipios_estoque_positivo":len(private),"vinculos_privados_elegiveis":sum(r.quality.get("vinculos_privados_elegiveis",0) for r in results),"codigos_nao_ligados":{},"reconciliacao_municipal_ok":sum(int(r["empregos_formais_privados"]) for r in private)==sum(r.quality.get("vinculos_privados_elegiveis",0) for r in results),"reconciliacao_territorial_ok":reconciliations,"semantica_base_nacional":"esparsa; ausência equivale a zero somente com coverage_complete=true"}
   self.interim.parent.mkdir(parents=True,exist_ok=True)
   with tempfile.TemporaryDirectory(prefix="sit-rais-publish-",dir=self.interim.parent) as d:
    stage=Path(d);_write_csv(stage/final_private.name,sorted(private,key=lambda r:r["municipio_id"]));_write_csv(stage/final_diag.name,sorted(diagnostic,key=lambda r:r["municipio_id"]));(stage/"qa.json").write_text(json.dumps(qa,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    self.interim.mkdir(parents=True,exist_ok=True);(stage/final_private.name).replace(final_private);(stage/final_diag.name).replace(final_diag);(stage/"qa.json").replace(qa_path)
   artifacts.extend(map(str,(final_private,final_diag,qa_path)));return CollectionResult(self.source,CollectionStatus.SUCCESS,artifacts,["MER-02","MER-DIAG-01"],entries,warnings)
  except ModuleNotFoundError as exc:return fail(CollectionStatus.BLOCKED_ENVIRONMENT,exc)
  except OfficialRoutesUnavailable as exc:entries.extend(exc.attempts);return fail(CollectionStatus.BLOCKED_SOURCE,exc)
  except (HTTPError,URLError,TimeoutError) as exc:return fail(CollectionStatus.BLOCKED_SOURCE,exc)
  except (ValueError,OSError,zipfile.BadZipFile) as exc:return fail(CollectionStatus.FAILED_VALIDATION,exc)
register_collector("rais",RaisCollector)
